import os
import sqlite3
import sys
import tempfile
import unittest
from datetime import datetime, timezone

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import leadfinder as lf  # noqa: E402


class KeywordTests(unittest.TestCase):
    def test_matches_whole_words_only(self):
        self.assertEqual(lf.keyword_hits("Need an inspection in Guangzhou", ["inspection"]), ["inspection"])

    def test_short_keyword_does_not_match_inside_words(self):
        self.assertEqual(lf.keyword_hits("the qcs were fine", ["QC"]), [])

    def test_case_insensitive(self):
        self.assertEqual(lf.keyword_hits("ALIBABA supplier", ["alibaba"]), ["alibaba"])


class ParseClassificationTests(unittest.TestCase):
    def test_plain_json(self):
        raw = ('{"is_lead": true, "lead_type": "inspection", "buyer_quality": 4, '
               '"urgency": 5, "product": "LED lamp", "summary": "x", "suggested_reply": "y"}')
        result = lf.parse_classification(raw)
        self.assertTrue(result["is_lead"])
        self.assertEqual(result["buyer_quality"], 4)
        self.assertEqual(result["urgency"], 5)

    def test_json_inside_code_fence(self):
        raw = 'Here you go:\n```json\n{"is_lead": false, "lead_type": "other"}\n```'
        result = lf.parse_classification(raw)
        self.assertFalse(result["is_lead"])

    def test_string_booleans_and_bad_values_are_normalized(self):
        raw = '{"is_lead": "false", "lead_type": "magic", "buyer_quality": 9, "urgency": "x"}'
        result = lf.parse_classification(raw)
        self.assertFalse(result["is_lead"])
        self.assertEqual(result["lead_type"], "other")
        self.assertEqual(result["buyer_quality"], 5)
        self.assertEqual(result["urgency"], 1)

    def test_unparseable_returns_none(self):
        self.assertIsNone(lf.parse_classification("sorry, I cannot help"))
        self.assertIsNone(lf.parse_classification('{"is_lead": true,'))


class DedupeTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.conn = lf.open_db(os.path.join(self.tmp.name, "seen.sqlite3"))

    def tearDown(self):
        self.conn.close()
        self.tmp.cleanup()

    def test_mark_then_seen(self):
        self.assertFalse(lf.is_seen(self.conn, "https://example.com/a"))
        lf.mark_seen(self.conn, "https://example.com/a", True)
        self.assertTrue(lf.is_seen(self.conn, "https://example.com/a"))

    def test_normalize_url_ignores_trailing_slash(self):
        lf.mark_seen(self.conn, lf.normalize_url("https://example.com/a/"), False)
        self.assertTrue(lf.is_seen(self.conn, lf.normalize_url("https://example.com/a")))


class XlsxTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = os.path.join(self.tmp.name, "leads.xlsx")

    def tearDown(self):
        self.tmp.cleanup()

    def row(self, url, buyer, urgency, status="new"):
        r = ["2026-10-17", "reddit", "r/Alibaba", "t", url, "u",
             "", "inspection", buyer, urgency, "s", "reply", status]
        return r

    def test_creates_header_and_sorts_by_score(self):
        lf.append_leads(self.path, [self.row("https://a", 2, 2), self.row("https://b", 5, 5),
                                    self.row("https://c", 3, 4)])
        import openpyxl
        ws = openpyxl.load_workbook(self.path).active
        self.assertEqual([c.value for c in ws[1]], lf.XLSX_HEADERS)
        urls = [ws.cell(row=i, column=lf.COL_URL + 1).value for i in range(2, ws.max_row + 1)]
        self.assertEqual(urls, ["https://b", "https://c", "https://a"])

    def test_keeps_existing_status_and_skips_duplicate_urls(self):
        lf.append_leads(self.path, [self.row("https://a", 3, 3)])
        import openpyxl
        wb = openpyxl.load_workbook(self.path)
        ws = wb.active
        ws.cell(row=2, column=lf.COL_STATUS + 1).value = "contacted"
        wb.save(self.path)

        added = lf.append_leads(self.path, [self.row("https://a", 3, 3), self.row("https://z", 1, 1)])
        self.assertEqual(added, 1)
        ws = openpyxl.load_workbook(self.path).active
        self.assertEqual(ws.max_row, 3)  # header + 2 leads
        self.assertEqual(ws.cell(row=2, column=lf.COL_STATUS + 1).value, "contacted")


class KeywordConfigTests(unittest.TestCase):
    def test_config_loads_and_has_required_sections(self):
        cfg = lf.load_config(os.path.join(os.path.dirname(__file__), "..", "config.yaml"))
        self.assertIn("subreddits", cfg["reddit"])
        self.assertIn("queries", cfg["web"])
        self.assertTrue(cfg["keywords"])


class FreshnessTests(unittest.TestCase):
    def test_window_maps_to_brave_codes(self):
        self.assertEqual(lf.freshness_for(1), "pd")
        self.assertEqual(lf.freshness_for(7), "pw")
        self.assertEqual(lf.freshness_for(30), "pm")
        self.assertEqual(lf.freshness_for(90), "py")


if __name__ == "__main__":
    unittest.main()
